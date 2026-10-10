"""World Bank Building Green dataset (building energy codes and standards, their requirements and enforcement, 88 economies;
DECIG, dataset version of 13 May 2024) as JSON for maps, every field kept, plus the original workbook unchanged.
Terms (World Bank, "Unlocking Efficiency" 2025, Rights and Permissions): "this work may be reproduced, in whole or in part,
for noncommercial purposes as long as full attribution to this work is given". So: noncommercial, attributed, kept apart.
  out/building-green.json  {source, licence_note, indicators: {code: name}, countries: {ISO3: {code: value}}}
  out/building-green.xlsx  the World Bank's workbook, unchanged
Usage: python3 bake.py"""
import datetime, json, os, urllib.request
import openpyxl
import sources

URL = 'https://www.worldbank.org/content/dam/sites/buildinggreen/doc/Building-green-dataset-final-05-13-2024.xlsx'
CITE = ('World Bank. 2025. Unlocking Efficiency: The Global Landscape of Building Energy Regulations and their Enforcement. '
        '© World Bank. Building Green dataset (https://www.worldbank.org/en/building-green).')


def main():
    os.makedirs(sources.OUT, exist_ok=True)
    xlsx = os.path.join(sources.OUT, 'building-green.xlsx')
    req = urllib.request.Request(URL, headers={'User-Agent': 'AlreadyOpen world-building-green (https://github.com/AlreadyOpen/world-building-green)'})
    open(xlsx, 'wb').write(urllib.request.urlopen(req, timeout=120).read())
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    names = {c: n for c, n in list(wb['Indicator Codes'].iter_rows(values_only=True))[1:] if c}
    rows = list(wb['Final dataset'].iter_rows(values_only=True))
    head, body = rows[0], [r for r in rows[1:] if r and r[1]]
    countries = {r[1]: {c: v for c, v in zip(head, r) if c} for r in body}
    out = {'source': 'World Bank, Global Indicators Group (DECIG): Building Green', 'url': 'https://www.worldbank.org/en/building-green',
           'dataset': URL, 'cite': CITE,
           'licence_note': 'Noncommercial reuse with full attribution only (World Bank, Unlocking Efficiency 2025, Rights and Permissions). '
                           'Not an open licence: do not use commercially.',
           'indicators': names, 'countries': countries}
    json.dump(out, open(os.path.join(sources.OUT, 'building-green.json'), 'w', encoding='utf-8'), ensure_ascii=False, default=str)
    idx = {'built': datetime.datetime.now().isoformat(timespec='seconds'),
           'counts': {'countries': len(countries), 'indicators': len(names)}}
    json.dump(idx, open(os.path.join(sources.OUT, 'index.json'), 'w'), indent=1)
    print(idx['counts'])


if __name__ == '__main__':
    main()
